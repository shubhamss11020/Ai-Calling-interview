import re
import os
from datetime import datetime
from typing import Dict, Any, List, Optional
import pypdf

import logging

logger = logging.getLogger(__name__)

SCANNED_FLAG = "SCANNED_DOCUMENT_REQUIRES_OCR"


class ResumeParserService:
    COMMON_SKILLS = [
        "python", "fastapi", "django", "flask", "postgresql", "mysql", "mongodb", "redis",
        "docker", "kubernetes", "aws", "gcp", "azure", "git", "ci/cd", "rest apis",
        "graphql", "celery", "node.js", "express", "react", "typescript", "javascript",
        "next.js", "tailwind", "linux", "microservices", "kafka", "rabbitmq", "pytorch",
        "tensorflow", "machine learning", "nlp", "llm", "langchain", "sql"
    ]

    MONTHS = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?"

    # ---- text extraction -------------------------------------------------

    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        """Extract plain text from PDF, tolerating unreadable pages."""
        text = ""
        try:
            reader = pypdf.PdfReader(file_path)
        except Exception as exc:
            logger.warning("Could not open PDF %s: %s", file_path, exc)
            return ""
        for page in reader.pages:
            try:
                extracted = page.extract_text()
            except Exception as exc:
                logger.warning("Skipping unreadable PDF page in %s: %s", file_path, exc)
                continue
            if extracted:
                text += extracted + "\n"
        return text

    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        try:
            import docx
            document = docx.Document(file_path)
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.extend(cell.text for cell in row.cells)
            return "\n".join(parts)
        except Exception as exc:
            logger.warning("Could not read DOCX %s: %s", file_path, exc)
            return ""

    @staticmethod
    def extract_text_from_txt(file_path: str) -> str:
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as fh:
                return fh.read()
        except Exception as exc:
            logger.warning("Could not read TXT %s: %s", file_path, exc)
            return ""

    @classmethod
    def extract_text(cls, file_path: str) -> str:
        """Dispatch on file extension (PDF, DOCX, TXT)."""
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".docx":
            return cls.extract_text_from_docx(file_path)
        if ext == ".txt":
            return cls.extract_text_from_txt(file_path)
        return cls.extract_text_from_pdf(file_path)

    # ---- field normalization ---------------------------------------------

    @staticmethod
    def normalize_phone(raw: Optional[str], default_country_code: str = "91") -> Optional[str]:
        """Normalize to E.164. Accepts +CC numbers and local 10-digit Indian numbers."""
        if not raw:
            return None
        has_plus = raw.strip().startswith("+")
        digits = re.sub(r"\D", "", raw)
        if has_plus and 8 <= len(digits) <= 15:
            return f"+{digits}"
        if len(digits) == 10:
            return f"+{default_country_code}{digits}"
        if len(digits) == 11 and digits.startswith("0"):
            return f"+{default_country_code}{digits[1:]}"
        if len(digits) == 12 and digits.startswith(default_country_code):
            return f"+{digits}"
        return None

    @classmethod
    def extract_phone(cls, text: str) -> Optional[str]:
        for match in re.finditer(r"(?:\+\d{1,3}[\s.-]?)?(?:\(?\d{3,5}\)?[\s.-]?)?\d{3,5}[\s.-]?\d{3,5}", text):
            normalized = cls.normalize_phone(match.group(0))
            if normalized:
                return normalized
        return None

    @classmethod
    def estimate_experience_years(cls, text: str) -> float:
        """Explicit '3.5 years' statements first, then employment date ranges, else 0.0."""
        clean = text.lower()
        explicit = [
            float(m.group(1))
            for m in re.finditer(r"(\d{1,2}(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)\b", clean)
            if float(m.group(1)) <= 50
        ]
        if explicit:
            return max(explicit)

        now_year = datetime.now().year
        range_re = re.compile(
            rf"(?:{cls.MONTHS}\s+)?((?:19|20)\d{{2}})\s*(?:-|–|—|to)\s*(?:{cls.MONTHS}\s+)?((?:19|20)\d{{2}}|present|current|now)"
        )
        spans = []
        for m in range_re.finditer(clean):
            start = int(m.group(1))
            end_raw = m.group(2)
            end = now_year if end_raw in ("present", "current", "now") else int(end_raw)
            if start <= end <= now_year:
                spans.append((start, end))
        if not spans:
            return 0.0
        # merge overlapping ranges so parallel roles aren't double counted
        spans.sort()
        merged = [list(spans[0])]
        for start, end in spans[1:]:
            if start <= merged[-1][1]:
                merged[-1][1] = max(merged[-1][1], end)
            else:
                merged.append([start, end])
        return float(sum(end - start for start, end in merged))

    # ---- full parse ------------------------------------------------------

    @classmethod
    def parse_resume_text(cls, text: str) -> Dict[str, Any]:
        """Extract candidate attributes, skills, and experience estimates."""
        clean_text = text.lower()

        email_match = re.search(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', text)
        email = email_match.group(0) if email_match else None

        found_skills: List[str] = []
        for skill in cls.COMMON_SKILLS:
            if re.search(r'(?<![\w.])' + re.escape(skill) + r'(?![\w])', clean_text):
                found_skills.append(skill.title() if not skill.isupper() else skill)

        is_scanned = len(text.strip()) < settings_min_chars()

        return {
            "email": email,
            "phone": cls.extract_phone(text),
            "extracted_skills": sorted(set(found_skills)),
            "total_experience_years": cls.estimate_experience_years(text),
            "raw_text": text,
            "parse_flag": SCANNED_FLAG if is_scanned else None,
        }


def settings_min_chars() -> int:
    from app.core.config import settings
    return settings.MIN_RESUME_TEXT_CHARS


resume_service = ResumeParserService()
